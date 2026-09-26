import io
from dataclasses import dataclass
from datetime import date
from decimal import Decimal as D

from app.application.importacao.mapeamento import ColunaMapeada, Mapeamento
from vitrine_core.datasets.tipos import TipoDataset

SETEMBRO = (date(2026, 9, 1), date(2026, 9, 30))
VENDEDORES = [
    ("MARIA SOUZA", 40, 110, D("8500.50"), D("300.00")),
    ("JOÃO LIMA", 55, 95, D("6200.00"), D("0.00")),
    ("ANA COSTA", 30, 88, D("4100.75"), D("150.25")),
]
TOTAL_BRUTO = sum(v[3] for v in VENDEDORES)


@dataclass(frozen=True)
class Layout:
    nome: str
    arquivo: str
    conteudo: bytes
    mapeamento: Mapeamento
    tem_total: bool
    vendedores_esperados: frozenset[str]
    bruto_esperado: D


def _mapeamento(cabecalho: int, colunas: dict[int, str], decimal: str = ",", periodo: bool = False) -> Mapeamento:
    return Mapeamento(
        tipo=TipoDataset.VENDAS_VENDEDOR_PERIODO,
        linha_cabecalho=cabecalho,
        colunas=[ColunaMapeada(indice=i, campo=c) for i, c in colunas.items()],
        separador_decimal=decimal,
        periodo={"inicio": SETEMBRO[0], "fim": SETEMBRO[1]} if periodo else None,
    )


def _br(valor: D) -> str:
    inteiro, centavos = f"{valor:.2f}".split(".")
    return f"{int(inteiro):,}".replace(",", ".") + "," + centavos


def _us(valor: D) -> str:
    return f"{valor:,.2f}"


def _layout(nome, arquivo, conteudo, mapeamento, tem_total, vendedores):
    return Layout(nome, arquivo, conteudo, mapeamento, tem_total, frozenset(vendedores), sum((v[3] for v in VENDEDORES), D(0)))


def csv_cp1252_com_codigo():
    linhas = ["Relatório de Vendas por Vendedor;;;;;", "Período: 01/09/2026 a 30/09/2026;;;;;", ";;;;;",
              "Cód;Nome do Vendedor;Cupons;Itens;Vlr Venda;Vlr Devolução"]
    linhas += [f"{i:03d};{n};{c};{p};{_br(b)};{_br(t)}" for i, (n, c, p, b, t) in enumerate(VENDEDORES, 1)]
    linhas.append(f"Totais;;{sum(v[1] for v in VENDEDORES)};{sum(v[2] for v in VENDEDORES)};{_br(TOTAL_BRUTO)};{_br(sum(v[4] for v in VENDEDORES))}")
    conteudo = ("\r\n".join(linhas) + "\r\n").encode("cp1252")
    return _layout("csv-cp1252-codigo", "vendas.csv", conteudo,
                   _mapeamento(3, {1: "vendedor", 2: "atendimentos", 3: "pecas", 4: "faturamento_bruto", 5: "trocas"}), True,
                   [v[0] for v in VENDEDORES])


def html_disfarcado_de_xls():
    linhas = "".join(
        f"<tr><td>{i:03d} - {n}</td><td>{c}</td><td>{p}</td><td>R$ {_br(b)}</td><td>{_br(t)}</td></tr>"
        for i, (n, c, p, b, t) in enumerate(VENDEDORES, 1)
    )
    html = f"""<html><head><meta charset="utf-8"></head><body>
    <table><tr><td colspan="5"><b>VENDAS POR CONSULTOR</b></td></tr>
    <tr><td colspan="5">Período: 01/09/2026 a 30/09/2026</td></tr>
    <tr><th>Consultor</th><th>Nº Vendas</th><th>Qtde Itens</th><th>Total Vendido</th><th>Estornos</th></tr>
    {linhas}
    <tr><td>Soma</td><td>{sum(v[1] for v in VENDEDORES)}</td><td>{sum(v[2] for v in VENDEDORES)}</td><td>R$ {_br(TOTAL_BRUTO)}</td><td>{_br(sum(v[4] for v in VENDEDORES))}</td></tr>
    </table></body></html>"""
    return _layout("html-xls-soma", "vendas.xls", html.encode("utf-8"),
                   _mapeamento(2, {0: "vendedor", 1: "atendimentos", 2: "pecas", 3: "faturamento_bruto", 4: "trocas"}), True,
                   [f"{i:03d} - {v[0]}" for i, v in enumerate(VENDEDORES, 1)])


def csv_utf16_tab_decimal_ponto():
    linhas = ["Seller\tTickets\tUnits\tAmount\tReturns"]
    linhas += [f"{n}\t{c}\t{p}\t{_us(b)}\t{_us(t)}" for n, c, p, b, t in VENDEDORES]
    conteudo = "\r\n".join(linhas).encode("utf-16")
    return _layout("csv-utf16-tab-ponto", "sales.txt", conteudo,
                   _mapeamento(0, {0: "vendedor", 1: "atendimentos", 2: "pecas", 3: "faturamento_bruto", 4: "trocas"}, decimal=".", periodo=True),
                   False, [v[0] for v in VENDEDORES])


def xlsx_numerico_sem_total():
    from openpyxl import Workbook

    livro = Workbook()
    planilha = livro.active
    planilha.append(["Vendedor", "Tickets", "Peças", "Faturamento", "Devoluções"])
    for n, c, p, b, t in VENDEDORES:
        planilha.append([n, c, p, float(b), float(t)])
    saida = io.BytesIO()
    livro.save(saida)
    return _layout("xlsx-numerico-sem-total", "vendas.xlsx", saida.getvalue(),
                   _mapeamento(0, {0: "vendedor", 1: "atendimentos", 2: "pecas", 3: "faturamento_bruto", 4: "trocas"}, periodo=True),
                   False, [v[0] for v in VENDEDORES])


def xlsx_varias_lojas_com_subtotais():
    from openpyxl import Workbook

    livro = Workbook()
    planilha = livro.active
    planilha.append(["RELATÓRIO CONSOLIDADO DE VENDEDORES"])
    planilha.append(["Período: 01/09/2026 a 30/09/2026"])
    planilha.append(["Vendedor", "Tickets", "Peças", "Valor", "Trocas"])
    lojas = {"LOJA CENTRO": VENDEDORES[:2], "LOJA NORTE": VENDEDORES[2:]}
    for loja, vendedores in lojas.items():
        planilha.append([loja])
        for n, c, p, b, t in vendedores:
            planilha.append([n, c, p, _br(b), _br(t)])
        planilha.append([f"SUBTOTAL {loja}", sum(v[1] for v in vendedores), sum(v[2] for v in vendedores),
                         _br(sum(v[3] for v in vendedores)), _br(sum(v[4] for v in vendedores))])
    planilha.append(["TOTAL", sum(v[1] for v in VENDEDORES), sum(v[2] for v in VENDEDORES), _br(TOTAL_BRUTO), _br(sum(v[4] for v in VENDEDORES))])
    saida = io.BytesIO()
    livro.save(saida)
    return _layout("xlsx-lojas-subtotais", "consolidado.xlsx", saida.getvalue(),
                   _mapeamento(2, {0: "vendedor", 1: "atendimentos", 2: "pecas", 3: "faturamento_bruto", 4: "trocas"}), True,
                   [v[0] for v in VENDEDORES])


def xls_legado():
    import xlwt

    livro = xlwt.Workbook()
    planilha = livro.add_sheet("Vendas")
    planilha.write(0, 0, "VENDAS POR VENDEDOR - SETEMBRO/2026")
    for c, titulo in enumerate(["Vendedor", "Qtd. Vendas", "Qtd. Peças", "Valor Bruto", "Devolvido"]):
        planilha.write(2, c, titulo)
    for r, (n, c, p, b, t) in enumerate(VENDEDORES, 3):
        for col, valor in enumerate([n, c, p, float(b), float(t)]):
            planilha.write(r, col, valor)
    linha = 3 + len(VENDEDORES)
    for col, valor in enumerate(["TOTAL GERAL", sum(v[1] for v in VENDEDORES), sum(v[2] for v in VENDEDORES), float(TOTAL_BRUTO), float(sum(v[4] for v in VENDEDORES))]):
        planilha.write(linha, col, valor)
    saida = io.BytesIO()
    livro.save(saida)
    return _layout("xls-legado", "vendas.xls", saida.getvalue(),
                   _mapeamento(2, {0: "vendedor", 1: "atendimentos", 2: "pecas", 3: "faturamento_bruto", 4: "trocas"}, periodo=True), True,
                   [v[0] for v in VENDEDORES])


def todos() -> list[Layout]:
    return [
        csv_cp1252_com_codigo(), html_disfarcado_de_xls(), csv_utf16_tab_decimal_ponto(),
        xlsx_numerico_sem_total(), xlsx_varias_lojas_com_subtotais(), xls_legado(),
    ]
