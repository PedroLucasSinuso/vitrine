import io
from datetime import date, datetime

VENDEDORES = [
    ("MARIANA SOUZA", 42, 118, "8.940,50", "320,00"),
    ("CARLOS LIMA", 55, 97, "6.215,30", "0,00"),
    ("JULIA PEREIRA", 38, 131, "7.402,90", "189,90"),
]
SUBTOTAL_CENTRO = ("SUBTOTAL LOJA CENTRO", 135, 346, "22.558,70", "509,90")
VENDEDORES_NORTE = [
    ("ANA COSTA", 20, 44, "3.100,00", "0,00"),
]
TOTAL_GERAL = ("TOTAL GERAL", 155, 390, "25.658,70", "509,90")


def _grade_vendedores(titulos_extras: int = 0, loja: str = "TACO SHOPPING CENTRO", total=TOTAL_GERAL):
    linhas = [
        ["RELATÓRIO DE VENDAS POR VENDEDOR", None, None, None, None, None],
        [f"Loja: {loja}", None, None, None, None, None],
    ]
    linhas += [[f"Filtro adicional {i}: todos", None, None, None, None, None] for i in range(titulos_extras)]
    linhas += [
        ["Período: 01/09/2026 a 30/09/2026", None, None, None, None, None],
        [],
        [None, "Vendas", None, None, None, None],
        ["Vendedor", "Qtd Tickets", "Qtd Peças", "Valor (R$)", "Trocas (R$)", None],
    ]
    linhas += [[n, t, p, f"R$ {v}", tr, None] for n, t, p, v, tr in VENDEDORES]
    linhas.append(list(SUBTOTAL_CENTRO) + [None])
    linhas += [[n, t, p, f"R$ {v}", tr, None] for n, t, p, v, tr in VENDEDORES_NORTE]
    linhas.append(list(total) + [None])
    linhas += [[], ["Emitido em 30/09/2026 18:04 por admin.loja", None, None, None, None, None]]
    return linhas


def linha_cabecalho_vendedores(titulos_extras: int = 0) -> int:
    return 5 + titulos_extras


def vendedores_xlsx(titulos_extras: int = 0, loja: str = "TACO SHOPPING CENTRO", total=TOTAL_GERAL) -> bytes:
    from openpyxl import Workbook

    livro = Workbook()
    planilha = livro.active
    for linha in _grade_vendedores(titulos_extras, loja, total):
        planilha.append(linha)
    cabecalho_superior = linha_cabecalho_vendedores(titulos_extras)
    planilha.merge_cells(start_row=cabecalho_superior, start_column=2, end_row=cabecalho_superior, end_column=4)
    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()


def vendedores_xls() -> bytes:
    import xlwt

    livro = xlwt.Workbook()
    planilha = livro.add_sheet("Relatorio")
    for r, linha in enumerate(_grade_vendedores()):
        for c, valor in enumerate(linha):
            if valor is not None:
                planilha.write(r, c, valor)
    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()


DIARIAS = [
    ("01/09/2026", "1.200,50", 14, 30),
    ("02/09/2026", "980,00", 11, 22),
    ("03/09/2026", "1.450,75", 17, 41),
]


def diarias_csv() -> bytes:
    linhas = [
        "Relatório diário de vendas;;;",
        "Emissão: 04/09/2026;;;",
        ";;;",
        "Data;Faturamento;Atendimentos;Peças",
    ]
    linhas += [f"{d};{v};{a};{p}" for d, v, a, p in DIARIAS]
    linhas.append("Total;3.631,25;42;93")
    return ("\r\n".join(linhas) + "\r\n").encode("cp1252")


def contatos_csv() -> bytes:
    linhas = [
        "vendedor,contatos,respostas,conversoes",
        "MARIANA SOUZA,120,48,9",
        "CARLOS LIMA,80,20,3",
    ]
    return "\n".join(linhas).encode("utf-8")


def itens_xlsx() -> bytes:
    from openpyxl import Workbook

    livro = Workbook()
    planilha = livro.active
    planilha.append(["Ticket", "Data", "Tipo", "Vendedor", "Produto", "Tamanho", "Cor", "Qtd", "Valor"])
    planilha.append(["1001", datetime(2026, 9, 1), "Venda", "MARIANA SOUZA", "Vestido Midi", "M", "Azul", 1, 259.9])
    planilha.append(["1001", datetime(2026, 9, 1), "Venda", "MARIANA SOUZA", "Cinto", "U", "Preto", 1, 79.9])
    planilha.append(["1002", datetime(2026, 9, 2), "Venda", "CARLOS LIMA", "Calça Jeans", "40", "Azul", 2, 399.8])
    planilha.append(["2001", datetime(2026, 9, 3), "Troca", "CARLOS LIMA", "Calça Jeans", "40", "Azul", 1, 199.9])
    planilha.append(["Total", None, None, None, None, None, None, 5, 939.5])
    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()


PRODUTOS = [("Vestido Midi", "FEMININO", 12, "3.118,80"), ("Calça Jeans", "FEMININO", 20, "3.998,00")]


def produtos_pdf(com_bordas: bool = True) -> bytes:
    from weasyprint import HTML

    estilo = "table, th, td { border: 1px solid #000; border-collapse: collapse; padding: 4px; }" if com_bordas else "td, th { padding: 4px 18px; }"

    linhas = "".join(f"<tr><td>{p}</td><td>{g}</td><td>{q}</td><td>{v}</td></tr>" for p, g, q, v in PRODUTOS)
    html = f"""
    <html><head><style>{estilo}</style></head><body>
      <h1>Curva de produtos</h1>
      <p>Período: 01/09/2026 a 30/09/2026</p>
      <table>
        <tr><th>Produto</th><th>Departamento</th><th>Quantidade</th><th>Receita</th></tr>
        {linhas}
        <tr><td>Total</td><td></td><td>32</td><td>7.116,80</td></tr>
      </table>
    </body></html>
    """
    return HTML(string=html).write_pdf()


def pdf_sem_tabela() -> bytes:
    from weasyprint import HTML

    return HTML(string="<p>Relatório sem tabela nenhuma</p>").write_pdf()


PERIODO_SETEMBRO = (date(2026, 9, 1), date(2026, 9, 30))
