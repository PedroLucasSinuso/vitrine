import csv
import io
import zipfile
from dataclasses import dataclass
from datetime import date, datetime

Celula = str | float | int | date | datetime | None

LIMITE_BYTES = 10 * 1024 * 1024
LIMITE_DESCOMPACTADO = 100 * 1024 * 1024
LIMITE_LINHAS = 50_000


class ArquivoInvalido(Exception):
    pass


@dataclass(frozen=True)
class Grade:
    formato: str
    linhas: list[list[Celula]]


def detectar_formato(conteudo: bytes, nome: str) -> str:
    if conteudo.startswith(b"%PDF"):
        return "pdf"
    if conteudo.startswith(b"\xd0\xcf\x11\xe0"):
        return "xls"
    if conteudo.startswith(b"PK"):
        return "xlsx"
    extensao = nome.rsplit(".", 1)[-1].lower() if "." in nome else ""
    if extensao in ("csv", "txt"):
        return "csv"
    raise ArquivoInvalido("Formato não suportado. Envie xls, xlsx, csv ou pdf.")


def _limpar(linhas: list[list[Celula]]) -> list[list[Celula]]:
    def vazia(c: Celula) -> bool:
        return c is None or (isinstance(c, str) and not c.strip())

    limpas = []
    for linha in linhas:
        linha = [c.strip() if isinstance(c, str) else c for c in linha]
        while linha and vazia(linha[-1]):
            linha.pop()
        limpas.append(linha)
    while limpas and not limpas[-1]:
        limpas.pop()
    if len(limpas) > LIMITE_LINHAS:
        raise ArquivoInvalido(f"O arquivo tem mais de {LIMITE_LINHAS} linhas.")
    return limpas


def _xlsx(conteudo: bytes) -> list[list[Celula]]:
    from openpyxl import load_workbook

    try:
        with zipfile.ZipFile(io.BytesIO(conteudo)) as pacote:
            if sum(i.file_size for i in pacote.infolist()) > LIMITE_DESCOMPACTADO:
                raise ArquivoInvalido("O arquivo descompactado é grande demais.")
        livro = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError) as erro:
        raise ArquivoInvalido("Não foi possível ler a planilha xlsx.") from erro
    try:
        planilha = next((p for p in livro.worksheets if p.max_row and p.max_row > 1), livro.active)
        return [list(linha) for linha in planilha.iter_rows(values_only=True)]
    finally:
        livro.close()


def _xls(conteudo: bytes) -> list[list[Celula]]:
    import xlrd

    try:
        livro = xlrd.open_workbook(file_contents=conteudo)
    except xlrd.XLRDError as erro:
        raise ArquivoInvalido("Não foi possível ler a planilha xls.") from erro
    planilha = next((livro.sheet_by_index(i) for i in range(livro.nsheets) if livro.sheet_by_index(i).nrows), livro.sheet_by_index(0))
    linhas = []
    for r in range(planilha.nrows):
        linha: list[Celula] = []
        for c in range(planilha.ncols):
            celula = planilha.cell(r, c)
            if celula.ctype == xlrd.XL_CELL_DATE:
                linha.append(xlrd.xldate.xldate_as_datetime(celula.value, livro.datemode))
            elif celula.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                linha.append(None)
            else:
                linha.append(celula.value)
        linhas.append(linha)
    return linhas


def _csv(conteudo: bytes) -> list[list[Celula]]:
    for codificacao in ("utf-8-sig", "cp1252"):
        try:
            texto = conteudo.decode(codificacao)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialeto = csv.Sniffer().sniff(texto[:5000], delimiters=";,\t|")
    except csv.Error:
        dialeto = csv.excel
    return [list(linha) for linha in csv.reader(io.StringIO(texto), dialeto)]


_PDF_POR_LINHAS = {"vertical_strategy": "lines", "horizontal_strategy": "lines"}
_PDF_POR_TEXTO = {"vertical_strategy": "text", "horizontal_strategy": "text"}


def _pdf(conteudo: bytes) -> list[list[Celula]]:
    import pdfplumber

    linhas: list[list[Celula]] = []
    try:
        with pdfplumber.open(io.BytesIO(conteudo)) as documento:
            for estrategia in (_PDF_POR_LINHAS, _PDF_POR_TEXTO):
                for pagina in documento.pages:
                    for tabela in pagina.extract_tables(estrategia):
                        linhas.extend([list(linha) for linha in tabela])
                if linhas:
                    break
    except Exception as erro:
        raise ArquivoInvalido("Não foi possível ler o PDF.") from erro
    if not linhas:
        raise ArquivoInvalido("Nenhuma tabela encontrada no PDF. PDF escaneado não é suportado.")
    return linhas


_LEITORES = {"xlsx": _xlsx, "xls": _xls, "csv": _csv, "pdf": _pdf}


def extrair(conteudo: bytes, nome: str) -> Grade:
    if len(conteudo) > LIMITE_BYTES:
        raise ArquivoInvalido("O arquivo passa de 10 MB.")
    if not conteudo:
        raise ArquivoInvalido("O arquivo está vazio.")
    formato = detectar_formato(conteudo, nome)
    linhas = _limpar(_LEITORES[formato](conteudo))
    if not linhas:
        raise ArquivoInvalido("O arquivo não tem dados.")
    return Grade(formato=formato, linhas=linhas)
